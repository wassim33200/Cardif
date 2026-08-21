"""End-to-end behaviour over the full generated corpus."""

import pandas as pd
import pytest
from openpyxl import load_workbook

from cardif.audit import write_audit
from cardif.pipeline import commit, profile_headers, run
from cardif.store import Warehouse


@pytest.fixture(scope="module")
def corpus_run(request):
    """One full pass over the corpus, reused by the tests in this module."""
    from cardif.config import load_config
    from pathlib import Path

    root, truth = request.getfixturevalue("corpus")
    config = load_config(Path(__file__).resolve().parent.parent / "config")
    # The one header the deterministic tiers cannot resolve; confirming it is what the
    # user does once in the review screen.
    config.aliases.learn("mtt glob", "prime_totale")
    outcome = run(root, config, "ade_immobilier", use_model=False)
    return root, truth, config, outcome


def test_every_workbook_is_processed(corpus_run):
    _, truth, _, outcome = corpus_run
    assert len(outcome.ok_results) == len(truth)
    assert outcome.skipped == []


def test_row_count_matches_the_answer_key(corpus_run):
    _, truth, _, outcome = corpus_run
    assert outcome.n_rows == sum(entry["n_rows"] for entry in truth)


def test_no_rows_vanish_unaccounted_for(corpus_run):
    """kept + dropped must equal what was in the source block. The property that
    guarantees cleaning never quietly loses a sale."""
    _, _, _, outcome = corpus_run
    for result in outcome.ok_results:
        table = result.table
        spanned = table.data_end - table.data_start + 1
        dropped_inside = [
            d for d in table.dropped_rows if table.data_start <= d.row <= table.data_end
        ]
        assert result.n_rows + len(dropped_inside) == spanned


def test_no_model_calls_are_needed_once_aliases_are_warm(corpus_run):
    _, _, _, outcome = corpus_run
    assert outcome.model_calls == 0
    assert outcome.unresolved_headers() == {}


def test_every_row_carries_provenance(corpus_run):
    _, _, _, outcome = corpus_run
    frame = outcome.frame
    for column in ("source_file", "source_sheet", "source_row", "row_hash"):
        assert frame[column].notna().all()
    assert frame["row_hash"].is_unique


def test_dates_are_a_real_date_type(corpus_run):
    """Object dtype here would cost the Power BI date-dimension relationship."""
    _, _, _, outcome = corpus_run
    assert pd.api.types.is_datetime64_any_dtype(outcome.frame["date_effet"])
    assert outcome.frame["date_effet"].notna().all()


def test_derived_premiums_are_labelled_not_disguised(corpus_run):
    """A computed total must never be indistinguishable from a reported one."""
    _, _, _, outcome = corpus_run
    frame = outcome.frame
    sources = set(frame["prime_totale_source"].dropna())
    assert sources <= {"reported", "derived"}

    derived = frame[frame["prime_totale_source"] == "derived"]
    assert not derived.empty
    # Where a total was derived, it must equal the parts it was derived from.
    assert (
        (derived["prime_nette"] + derived["frais"] - derived["prime_totale"]).abs() < 0.01
    ).all()


def test_late_declarations_are_kept_not_dropped(corpus_run):
    """The reason date_effet and mois_reception are separate columns."""
    _, _, _, outcome = corpus_run
    frame = outcome.frame
    effect_month = frame["date_effet"].dt.strftime("%Y-%m")
    assert (effect_month != frame["mois_reception"]).any()


def test_run_is_deterministic(corpus_run):
    """Same inputs, same output. Without this an auditor cannot check anything."""
    root, _, config, outcome = corpus_run
    again = run(root, config, "ade_immobilier", use_model=False)
    assert again.frame["row_hash"].tolist() == outcome.frame["row_hash"].tolist()
    first = outcome.frame.drop(columns=["ingested_at"])
    second = again.frame.drop(columns=["ingested_at"])
    pd.testing.assert_frame_equal(first, second)


def test_profile_headers_inventories_every_variant(corpus_run):
    root, _, config, _ = corpus_run
    inventory = profile_headers(root, config)
    assert not inventory.empty
    assert {"variantes", "banques", "fichiers", "champ_propose"} <= set(inventory.columns)
    # The inventory is what the target schema is chosen from, so it must show the
    # different wordings banks use for one field side by side.
    agence = inventory[inventory["champ_propose"] == "agence"]
    assert len(agence) > 1


def test_commit_then_recommit_is_idempotent(corpus_run, tmp_path):
    root, _, config, outcome = corpus_run
    warehouse_path = tmp_path / "warehouse"

    first = commit(outcome, config, warehouse_path)
    second = commit(outcome, config, warehouse_path)

    assert first["total"] == second["total"] == outcome.n_rows
    assert len(Warehouse(config, warehouse_path).read_fact()) == outcome.n_rows


def test_audit_workbook_has_every_sheet(corpus_run, tmp_path):
    _, _, _, outcome = corpus_run
    path = write_audit(
        tmp_path / "audit.xlsx", outcome.results, outcome.validation, outcome.frame
    )
    workbook = load_workbook(path)
    assert {"Résumé", "Correspondances", "Supprimé", "Anomalies", "Rapprochement"} <= set(
        workbook.sheetnames
    )
    # Every dropped row must be findable in the source from the audit alone.
    dropped = workbook["Supprimé"]
    headers = [c.value for c in dropped[1]]
    assert {"fichier", "position", "motif"} <= set(headers)


def test_reconciliation_totals_match_the_fact_table(corpus_run):
    from cardif.validate import reconciliation_table

    _, _, _, outcome = corpus_run
    table = reconciliation_table(outcome.frame)
    assert table["lignes"].sum() == outcome.n_rows
    assert table["prime_totale_total"].sum() == pytest.approx(
        outcome.frame["prime_totale"].sum(), abs=0.5
    )
