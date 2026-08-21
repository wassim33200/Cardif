"""The sample data and the things a non-technical colleague sees.

The sample exists to be someone's first impression of the tool. If it ever fails, or
produces warnings, or looks nothing like a real bank file, it does more harm than not
existing. So it is tested like a feature, not like a fixture.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

RACINE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RACINE / "app"))

from cardif.demo import creer_donnees_exemple      # noqa: E402
from cardif.pipeline import commit, run            # noqa: E402
from cardif.store import Warehouse                 # noqa: E402


@pytest.fixture(scope="module")
def exemple(tmp_path_factory) -> tuple[Path, dict]:
    dossier = tmp_path_factory.mktemp("exemple")
    resume = creer_donnees_exemple(dossier)
    return dossier, resume


class TestLesFichiersProduits:
    def test_un_dossier_par_banque(self, exemple):
        dossier, resume = exemple
        dossiers = [p for p in dossier.iterdir() if p.is_dir()]
        assert len(dossiers) == resume["banques"]
        assert all(any(p.glob("*.xlsx")) for p in dossiers)

    def test_douze_mois_par_produit_et_par_banque(self, exemple):
        """Chaque banque envoie un fichier par produit et par mois."""
        dossier, resume = exemple
        total = sum(
            len(list(banque.glob("*.xlsx")))
            for banque in dossier.iterdir() if banque.is_dir()
        )
        assert total == resume["fichiers"]
        assert total % 12 == 0, "chaque produit doit couvrir douze mois"

    def test_plusieurs_produits_par_banque(self, exemple):
        """Le sujet même de l'outil : une banque vend plusieurs produits."""
        import sys
        sys.path.insert(0, str(RACINE / "src"))
        from cardif.config import load_config
        from cardif.filemeta import scan

        dossier, resume = exemple
        config = load_config(RACINE / "config")
        metas = scan(dossier, config.banks, config.schema_)

        par_banque: dict[str, set[str]] = {}
        for meta in metas:
            if meta.bank_code and meta.produit:
                par_banque.setdefault(meta.bank_code, set()).add(meta.produit)

        assert par_banque, "aucun fichier rattaché à une banque et à un produit"
        for banque, produits in par_banque.items():
            assert len(produits) > 1, f"{banque} ne vend qu'un produit dans l'exemple"

    def test_chaque_fichier_nomme_son_produit(self, exemple):
        """Sans cela, l'outil ne saurait pas quel format appliquer."""
        import sys
        sys.path.insert(0, str(RACINE / "src"))
        from cardif.config import load_config
        from cardif.filemeta import scan

        dossier, _ = exemple
        config = load_config(RACINE / "config")
        non_identifies = [
            m.path.name for m in scan(dossier, config.banks, config.schema_)
            if m.produit is None
        ]
        assert non_identifies == []

    def test_une_note_explique_ce_que_sont_ces_fichiers(self, exemple):
        dossier, _ = exemple
        note = dossier / "LISEZ-MOI.txt"
        assert note.exists()
        texte = note.read_text(encoding="utf-8")
        assert "Aucune donnée réelle" in texte

    def test_les_fichiers_ont_bien_les_defauts_annonces(self, exemple):
        """A sample that is clean would teach the wrong thing about the tool."""
        dossier, _ = exemple
        classeurs = sorted(dossier.rglob("*.xlsx"))

        decale = 0
        avec_total = 0
        for chemin in classeurs:
            classeur = load_workbook(chemin, data_only=True)
            feuille = classeur.active
            premiere_remplie = next(
                (ligne for ligne in range(1, 20)
                 if any(feuille.cell(row=ligne, column=c).value is not None
                        for c in range(1, 10))),
                1,
            )
            if premiere_remplie > 1:
                decale += 1
            valeurs = [
                str(feuille.cell(row=ligne, column=1).value or "").strip().upper()
                for ligne in range(1, feuille.max_row + 1)
            ]
            if "TOTAL" in valeurs:
                avec_total += 1
            classeur.close()

        assert decale > 0, "aucun fichier ne commence après la première ligne"
        assert avec_total > 0, "aucun fichier ne porte de ligne TOTAL"

    def test_les_noms_de_colonnes_varient_entre_les_mois(self, exemple):
        """The whole point of the column step is that banks are inconsistent."""
        dossier, _ = exemple
        banque = next(p for p in dossier.iterdir() if p.is_dir())
        entetes = set()
        for chemin in banque.glob("*.xlsx"):
            classeur = load_workbook(chemin, data_only=True)
            feuille = classeur.active
            for ligne in feuille.iter_rows(min_row=1, max_row=15):
                textes = [str(c.value) for c in ligne if isinstance(c.value, str)]
                if len(textes) >= 4:
                    entetes.update(textes)
                    break
            classeur.close()
        assert len(entetes) > 8, "les en-têtes ne varient pas d'un mois à l'autre"


class TestLeParcoursDeDemonstration:
    """What a colleague will actually see when they click the button."""

    def test_tout_passe_sans_une_seule_erreur(self, exemple, config):
        dossier, resume = exemple
        resultat = run(dossier, config, use_model=False)

        assert resultat.skipped == [], "un fichier d'exemple n'a pas pu être traité"
        assert len(resultat.ok_results) == resume["fichiers"]
        assert resultat.n_rows == resume["ventes"]

    def test_aucune_colonne_ne_reste_a_deviner(self, exemple, config):
        """The demo must not open on a question the colleague cannot answer."""
        dossier, _ = exemple
        resultat = run(dossier, config, use_model=False)
        assert resultat.unresolved_headers() == {}

    def test_aucune_anomalie_bloquante(self, exemple, config):
        dossier, _ = exemple
        resultat = run(dossier, config, use_model=False)
        assert resultat.validation.errors == 0

    def test_aucun_appel_au_modele_n_est_necessaire(self, exemple, config):
        """The sample has to work with nothing else installed."""
        dossier, _ = exemple
        assert run(dossier, config, use_model=False).model_calls == 0

    def test_la_base_est_ecrite_et_relisible(self, exemple, config, tmp_path):
        dossier, resume = exemple
        entrepot = tmp_path / "base"
        resultat = run(dossier, config, use_model=False)
        resume_ecriture = commit(resultat, config, entrepot)

        assert resume_ecriture["total"] == resume["ventes"]
        # Une table par produit, pas un fichier unique.
        tables = sorted(entrepot.glob("fact_*.parquet"))
        assert len(tables) > 1, "un seul fichier : les produits n'ont pas été séparés"
        assert (entrepot / "dim_date.parquet").exists()
        assert len(Warehouse(config, entrepot).read_fact()) == resume["ventes"]

    def test_les_deux_formes_de_prime_sont_representees(self, exemple, config):
        """One bank reports a breakdown, another only the global figure."""
        dossier, _ = exemple
        frame = run(dossier, config, use_model=False).frame
        assert frame["prime_totale"].notna().all()
        assert frame["banque"].nunique() >= 2

    def test_les_montants_ecrits_en_texte_sont_bien_lus(self, exemple, config):
        """One bank writes "1 234,56" rather than a number."""
        dossier, _ = exemple
        frame = run(dossier, config, use_model=False).frame
        assert pd.api.types.is_numeric_dtype(frame["prime_totale"])
        assert frame["prime_totale"].min() > 0


class TestAffichageLisible:
    """Nothing on screen may be a machine value or an internal column name."""

    PRODUIT = "sahti"

    def _frame(self):
        return pd.DataFrame({
            "banque": ["CNEP"], "produit_code": ["sahti"],
            "mois_reception": ["2025-03"], "num_contrat": ["C-1"],
            "nom_client": ["A B"], "date_naissance": pd.to_datetime(["1980-01-01"]),
            "agence": ["Alger Centre"], "date_effet": pd.to_datetime(["2025-07-16"]),
            "date_fin": pd.to_datetime(["2026-07-15"]), "formule": ["Familiale"],
            "nb_assures": [3], "capital_assure": [442988.76],
            "periodicite": ["Mensuel"], "prime_totale": [3519.42],
            "prime_totale_source": ["derived"], "commission_partenaire": [120.0],
            "statut": ["En cours"],
            "source_file": ["x.xlsx"], "source_sheet": ["F"], "source_row": [7],
            "ingested_at": ["t"], "row_hash": ["h"],
        })

    def test_les_colonnes_techniques_sont_masquees(self, config):
        from composants import pour_affichage

        colonnes = list(pour_affichage(self._frame(), config, self.PRODUIT).columns)
        for interne in ("source_file", "row_hash", "ingested_at", "source_row"):
            assert interne not in colonnes

    def test_elles_restent_accessibles_a_la_demande(self, config):
        from composants import pour_affichage

        colonnes = list(
            pour_affichage(self._frame(), config, self.PRODUIT, technique=True).columns
        )
        assert "Fichier d'origine" in colonnes

    def test_les_dates_perdent_l_heure(self, config):
        from composants import pour_affichage

        affiche = pour_affichage(self._frame(), config, self.PRODUIT)
        assert affiche["Date d'effet"].iloc[0] == "16/07/2025"

    def test_les_valeurs_machine_deviennent_des_mots(self, config):
        from composants import pour_affichage

        affiche = pour_affichage(self._frame(), config, self.PRODUIT)
        assert affiche["Origine prime"].iloc[0] == "calculée"

    def test_aucun_nom_de_colonne_interne_ne_reste(self, config):
        from composants import pour_affichage

        for colonne in pour_affichage(self._frame(), config, self.PRODUIT).columns:
            assert "_" not in colonne, f"nom interne affiché : {colonne}"


class TestLancement:
    """The colleague never opens a terminal, so the launchers must be present."""

    @pytest.mark.parametrize(
        "nom",
        ["Installer.bat", "Lancer Cardif.bat", "installer.command", "lancer-cardif.command"],
    )
    def test_les_raccourcis_existent(self, nom):
        assert (RACINE / nom).exists()

    @pytest.mark.parametrize("nom", ["Lancer Cardif.bat", "lancer-cardif.command"])
    def test_ils_pointent_vers_le_bon_fichier(self, nom):
        contenu = (RACINE / nom).read_text(encoding="utf-8", errors="ignore")
        assert "Accueil.py" in contenu

    def test_les_raccourcis_mac_sont_executables(self):
        import os
        import stat

        for nom in ("installer.command", "lancer-cardif.command"):
            mode = (RACINE / nom).stat().st_mode
            assert mode & stat.S_IXUSR, f"{nom} n'est pas exécutable"

    def test_streamlit_ne_transmet_aucune_statistique(self):
        reglages = (RACINE / ".streamlit" / "config.toml").read_text(encoding="utf-8")
        assert "gatherUsageStats = false" in reglages
        assert 'address = "127.0.0.1"' in reglages

    def test_un_guide_existe_pour_les_collegues(self):
        guide = (RACINE / "GUIDE.md").read_text(encoding="utf-8")
        assert "Lancer Cardif" in guide
        assert "Installer.bat" in guide


class TestAffichageDesProduits:
    """What a colleague reads on screen must never be an internal value."""

    def _frame_produit(self):
        return pd.DataFrame({
            "banque": ["CNEP"], "produit_code": ["sahti"],
            "mois_reception": ["2025-03"], "num_contrat": ["C-1"],
            "nom_client": ["A B"], "date_naissance": [pd.NaT],
            "agence": ["Alger Centre"], "date_effet": pd.to_datetime(["2025-07-16"]),
            "date_fin": [pd.NaT], "formule": ["Familiale"], "nb_assures": [3],
            "capital_assure": [442988.76], "periodicite": ["Mensuel"],
            "prime_totale": [3519.42], "commission_partenaire": [120.0],
            "statut": ["En cours"], "source_file": ["x.xlsx"], "source_sheet": ["F"],
            "source_row": [7], "ingested_at": ["t"], "row_hash": ["h"],
        })

    def test_le_code_produit_devient_son_nom(self, config):
        from composants import pour_affichage

        affiche = pour_affichage(self._frame_produit(), config, "sahti")
        assert affiche["Code produit"].iloc[0] == "SAHTI"

    def test_une_case_vide_reste_vide(self, config):
        """« None » à l'écran ressemble à une valeur ; ce n'en est pas une."""
        from composants import pour_affichage

        affiche = pour_affichage(self._frame_produit(), config, "sahti")
        assert affiche["Date de fin"].iloc[0] == ""
        assert "None" not in affiche.astype(str).to_string()

    def test_chaque_produit_saffiche_avec_ses_propres_colonnes(self, config):
        from composants import pour_affichage

        colonnes = list(pour_affichage(self._frame_produit(), config, "sahti").columns)
        assert "Nombre d'assurés" in colonnes
        assert "Capital restant dû" not in colonnes


class TestExempleMultiProduits:
    def test_les_agences_sont_algeriennes(self, exemple):
        """Le partenaire est algérien : des agences tunisiennes sèmeraient le doute."""
        dossier, _ = exemple
        classeur = load_workbook(sorted(dossier.rglob("*.xlsx"))[0], data_only=True)
        feuille = classeur.active
        textes = {
            str(c.value) for ligne in feuille.iter_rows() for c in ligne
            if isinstance(c.value, str)
        }
        assert not {"Tunis Centre", "Sfax Nord", "Sousse Medina"} & textes

    def test_le_resume_annonce_les_produits(self, exemple):
        _, resume = exemple
        assert resume["produits"] > 1
        assert resume["fichiers"] == resume["produits"] * 12 or resume["fichiers"] % 12 == 0
