"""Regression cases transcribed from the supplied folder tree and ADE dictionary."""
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from cardif.filemeta import read_meta, resolve_period, scan
from cardif.mapping import Mapper
from cardif.pipeline import commit, profile_headers, run


@pytest.mark.parametrize('folder,product', [
    ('ADE CONSO', 'ade_consommation'),
    ('ADE IMMO', 'ade_immobilier'),
    ('ASSUR COMPTE', 'assurcompte'),
    ('CARTE VISA', 'voyage_visa'),
    ('Prévoyance Collective', 'prevoyance_collective'),
    ('Prévoyance individuelle', 'prevoyance_individuelle'),
    ('Prévoyance indiciduelle', 'prevoyance_individuelle'),
])
def test_screenshot_folder_layout(tmp_path, config, folder, product):
    path = tmp_path / 'BNP' / folder / '03-2026' / 'ventes.xlsx'
    path.parent.mkdir(parents=True)
    path.touch()
    meta, = scan(tmp_path, config.banks, config.schema_)
    assert (meta.bank_code, meta.produit, meta.period) == ('BNPPED', product, '2026-03')
    assert meta.period_method == 'folder_numeric'
    assert meta.resolved


@pytest.mark.parametrize('name,expected', [
    ('P_ADE_012026_T24.xlsx', (2026, 1)),
    ('C_ADE_122026_D6.xlsx', (2026, 12)),
    ('P_ADE_202603-FIIT.xlsx', (2026, 3)),
    ('Imp_ADE_IMMO_032026.xlsx', (2026, 3)),
    ('P_ADE_132026_T24.xlsx', (None, None)),
])
def test_compact_dates(name, expected):
    year, month, _, _ = resolve_period(name)
    assert (year, month) == expected


def test_conflicting_month_is_blocked_before_open(config):
    path = Path('BNP/ADE IMMO/02-2026/P_ADE_032026_T24.xlsx')
    meta = read_meta(path, config.banks, config.schema_)
    assert meta.period_conflict
    assert not meta.resolved
    from cardif.consolidate import process_file
    result = process_file(meta, config, Mapper(config))
    assert not result.ok
    assert 'Périodes contradictoires' in result.error


# Explicit samples: expected destinations below do not derive from the config.
SAMPLES = {
    'P_ADE_032026_T24.xlsx': (
        ['LD', 'AGENCE', 'CLIENT', 'NOM_CLIENT', 'PRENOM_CLIENT',
         'DATE_MOBILISATION', 'PRIME', 'MOIS DE RECEPTION'],
        ['0001', '001', 'C001', 'Nom Test', 'Prenom Test',
         '01/03/2026', '1 200,50', '02-2026']),
    'P_ADE_032026_D6.xlsx': (
        ['NUM_POLICE', 'N_PRET', 'NUM_CONTRA', 'CODE_AGENC', 'NUM_PART',
         'NOM', 'PRENOM', 'DATE_EFFET', 'MNT_PRIME', 'DFIN_CREDI', 'DFIN_ASSUR'],
        ['0002', '000020', '000030', '001', 'C002', 'Nom Test', 'Prenom Test',
         '02/03/2026', 1300, '01/03/2030', '01/03/2031']),
    'C_ADE_032026_T24.xlsx': (
        ['N_PRET', 'CODE_AGENC', 'DATE_MOBILISATION', 'Mnt PRIMES G',
         'MNT_PRIME', 'TYPE_CREDIT', 'STATUT_PRIME', 'MOIS PRIME', 'ANNEE PRIME'],
        ['0003', '001', '03/03/2026', 1400, 70, 'ASSURANCE-X', 'TYPE-X', '02', '2026']),
    'C_ADE_032026_D6.xlsx': (
        ['NUM_POLICE', 'N_PRET', 'IDAGENCE', 'AGENCE', 'Prime', 'MNT_PRIME',
         'MOIS', 'ANNEE', 'DATE_VAL', 'RIB'],
        ['0004', '000040', '001', 'Agence Technique', 1500, 80,
         '02', '2026', '04/03/2026', '0000123456789']),
    'P_ADE_032026-FIIT.xlsx': (
        ['N_PRET', 'CODE_AGENC', 'DATE_MOBILISATION', 'PRIME TOTALE', 'MNT_PRIME',
         'TAUX_SUBPRIME2', 'SUBPRIME2', 'MOIS', 'Année', 'CODE PRODUIT', 'CODE_C'],
        ['0005', '001', '05/03/2026', 1600, 90, '2%', 11, '02', '2026', 'P01', 'C01']),
    'C_ADE_032026-FIT.xlsx': (
        ['N_PRET', 'CODE_AGENC', 'DATE_MOBILISATION', 'PRIME TOTALE',
         'MNT_PRIME', 'MOIS PRIME', 'ANNEE PRIME', 'MOIS', 'Année'],
        ['0006', '001', '06/03/2026', 1700, 100, '01', '2025', '02', '2026']),
    'Imp_ADE_IMMO_032026.xlsx': (
        ['NUM_POLICE', 'NUM_PRET', 'IDAGENCE', 'MNT_ASSUR', 'NUM_PART',
         'DATE_EFFET', 'DFIN_ASSUR', 'MNT_PRIME', 'DATE_EFFET_THEO', 'REMARQUE'],
        ['0007', '000070', '001', 500000, 'C007', '07/03/2026', '01/03/2030',
         1800, '01/02/2026', 'Conserver cette remarque']),
}


@pytest.fixture
def actual_files(tmp_path):
    root = tmp_path / 'inputs'
    folder = root / 'CNEP' / 'ADE IMMO' / '03-2026'
    folder.mkdir(parents=True)
    for name, (headers, values) in SAMPLES.items():
        wb = Workbook()
        ws = wb.active
        ws.append(['Export mensuel'])
        ws.append([])
        ws.append(headers)
        for i in range(3):
            row = list(values)
            row[0] += str(i)
            ws.append(row)
        wb.save(folder / name)
        wb.close()
    return root


def test_seven_formats_keep_distinct_dictionary_fields(actual_files, config, tmp_path):
    outcome = run(actual_files, config, use_model=False)
    assert not outcome.failed_results
    assert outcome.n_rows == 21
    assert outcome.model_calls == 0
    assert not outcome.validation.errors, outcome.validation.to_frame().to_string()
    assert not outcome.unresolved_headers()
    frames = {r.meta.path.name: r.frame for r in outcome.ok_results}
    pt = frames['P_ADE_032026_T24.xlsx'].iloc[0]
    assert pt['num_contrat'] == '00010'
    assert pt['ade_client'] == 'C001'
    assert pt['nom_client'] == 'Nom Test'
    assert pt['ade_prenom_client'] == 'Prenom Test'
    assert pt['prime_totale'] == 1200.5
    assert pt['mois_reception'] == '2026-03'
    assert pt['ade_mois_de_reception'] == '02-2026'
    pd6 = frames['P_ADE_032026_D6.xlsx'].iloc[0]
    assert pd6['num_credit'] == '000020'
    assert pd6['ade_num_contra'] == '000030'
    assert pd6['date_fin'] == pd.Timestamp('2030-03-01')
    assert pd6['ade_dfin_assur'] == pd.Timestamp('2031-03-01')
    ct = frames['C_ADE_032026_T24.xlsx'].iloc[0]
    assert (ct['prime_totale'], ct['ade_mnt_prime']) == (1400, 70)
    assert ct['ade_assurance'] == 'ASSURANCE-X'
    assert ct['ade_type_assurance'] == 'TYPE-X'
    cd = frames['C_ADE_032026_D6.xlsx'].iloc[0]
    assert (cd['agence'], cd['ade_agence_']) == ('001', 'Agence Technique')
    assert (cd['prime_totale'], cd['ade_mnt_prime']) == (1500, 80)
    assert cd['ade_rib'] == '0000123456789'
    assert pd.isna(cd['date_effet'])  # Not reported in this collection format.
    fit = frames['P_ADE_032026-FIIT.xlsx'].iloc[0]
    assert (fit['ade_code_ass'], fit['ade_code_c']) == ('P01', 'C01')
    assert fit['ade_taux_prime_sub'] == 0.02
    assert pd.isna(fit['ade_taux_subprime2'])
    imp = frames['Imp_ADE_IMMO_032026.xlsx'].iloc[0]
    assert imp['montant_credit'] == 500000
    assert pd.isna(imp['ade_mnt_assur'])
    assert imp['ade_remarque'] == 'Conserver cette remarque'
    assert imp['ade_date_effet_theo'] == pd.Timestamp('2026-02-01')

    target = tmp_path / 'warehouse'
    saved = commit(outcome, config, target)
    assert saved['written'] == 21
    fact = pd.read_parquet(target / 'fact_ade_immobilier.parquet')
    assert len(fact) == 21
    workbook = load_workbook(target / 'par_banque/CNEP_ade_immobilier.xlsx', read_only=True)
    try:
        headers = next(workbook.active.values)
        assert headers[:5] == ('LD', 'AGENCE', 'MONTANT_ACCORDE', 'CRD', 'CLIENT')
        assert headers[117:120] == ('DATE_EFFET_THEO', 'DOS_TYPE', 'REMARQUE')
        assert len(headers) == len(set(headers))
    finally:
        workbook.close()


def test_profile_keeps_source_specific_meanings(actual_files, config):
    inventory = profile_headers(actual_files, config)
    matches = inventory[inventory['en_tete_normalise'] == 'n pret']
    assert set(matches['champ_propose']) == {'num_contrat', 'num_credit'}
    assert matches['format_source'].nunique() == 5


def test_unknown_headers_are_not_guessed(config):
    mapper = Mapper(config, lambda *args: pytest.fail('Must not guess dictionary fields'))
    source = config.source_format('C_ADE_032026_D6.xlsx', 'ade_immobilier')
    unknown = mapper.resolve_header('MNT_PRIMEE', 0, [1, 2], source_format=source)
    assert unknown.canonical is None
    # Source dictionary cannot change an unrelated product's meaning of CLIENT.
    assert mapper.resolve_header('CLIENT', 0, ['Nom Test'], produit='sahti').canonical == 'nom_client'
    assert config.source_format('C_ADE_032026_D6.xlsx', 'sahti') is None


def test_same_format_resend_still_flags_duplicates(actual_files, config):
    folder = actual_files / 'CNEP/ADE IMMO/03-2026'
    target = folder / 'copie/P_ADE_032026_T24.xlsx'
    target.parent.mkdir()
    target.write_bytes((folder / 'P_ADE_032026_T24.xlsx').read_bytes())
    outcome = run(actual_files, config, use_model=False)
    assert any(f.code == 'period_covered_twice' for f in outcome.validation.flags)


def test_unknown_format_does_not_hide_a_resend():
    from cardif.validate import check_period_coverage
    frame = pd.DataFrame({
        'banque': ['CNEP', 'CNEP'], 'mois_reception': ['2026-03', '2026-03'],
        'source_file': ['P_ADE_032026_T24.xlsx', 'P_ADE_032026_T24_corrige.xlsx'],
        'source_format': ['p_t24', None],
    })
    assert len(check_period_coverage(frame)) == 1
