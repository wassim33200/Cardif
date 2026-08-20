"""End-to-end runs, shared by the CLI and the Streamlit app.

Keeping the orchestration here means the UI and the command line cannot drift apart:
both call the same functions and see the same results.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

from .config import Config
from .consolidate import FileResult, process_file
from .filemeta import FileMeta, scan
from .llm import LocalModel, make_resolver
from .mapping import Mapper
from .store import Warehouse
from .validate import ValidationReport, validate


@dataclass
class RunResult:
    """Everything one pass over a data root produced."""

    results: list[FileResult] = field(default_factory=list)
    profile_name: str = ""
    frame: pd.DataFrame | None = None
    validation: ValidationReport = field(default_factory=ValidationReport)
    skipped: list[tuple[Path, str]] = field(default_factory=list)
    model_calls: int = 0

    @property
    def ok_results(self) -> list[FileResult]:
        return [r for r in self.results if r.ok]

    @property
    def failed_results(self) -> list[FileResult]:
        return [r for r in self.results if not r.ok]

    @property
    def n_rows(self) -> int:
        return 0 if self.frame is None else len(self.frame)

    def unresolved_headers(self) -> Counter:
        """Headers still needing a human, counted across files."""
        counter: Counter = Counter()
        for result in self.results:
            for mapping in result.unresolved_headers:
                counter[mapping.header] += 1
        return counter

    def summary_line(self) -> str:
        parts = [
            f"{len(self.ok_results)}/{len(self.results)} files",
            f"{self.n_rows} rows",
        ]
        if self.validation.errors:
            parts.append(f"{self.validation.errors} errors")
        if self.validation.warnings:
            parts.append(f"{self.validation.warnings} warnings")
        if self.model_calls:
            parts.append(f"{self.model_calls} model calls")
        return ", ".join(parts)


def build_mapper(config: Config, use_model: bool = True) -> tuple[Mapper, LocalModel | None]:
    """Construct the mapper, attaching the local model only if it is actually reachable.

    An unreachable endpoint degrades to the deterministic tiers rather than failing the
    run: the tool must stay usable when LM Studio simply is not running.
    """
    if not use_model or not config.settings.llm.enabled:
        return Mapper(config), None
    try:
        model = LocalModel(config.settings.llm)
    except Exception:
        return Mapper(config), None
    if not model.available():
        return Mapper(config), None
    return Mapper(config, make_resolver(config, model)), model


def run(
    root: Path | str,
    config: Config,
    profile_name: str | None = None,
    use_model: bool = True,
    only: list[Path] | None = None,
    overrides: dict[str, str] | None = None,
) -> RunResult:
    """Process every workbook under ``root`` (or just ``only``) into one frame."""
    profile_name = profile_name or config.settings.warehouse.profile
    mapper, model = build_mapper(config, use_model)

    metas: list[FileMeta] = scan(Path(root), config.banks)
    if only is not None:
        wanted = {str(p) for p in only}
        metas = [m for m in metas if str(m.path) in wanted]

    outcome = RunResult(profile_name=profile_name)
    frames = []
    for meta in metas:
        result = process_file(meta, config, mapper, profile_name, overrides)
        outcome.results.append(result)
        if result.ok:
            frames.append(result.frame)
        else:
            outcome.skipped.append((meta.path, result.error or "no rows"))

    if frames:
        outcome.frame = pd.concat(frames, ignore_index=True)
        outcome.validation = validate(outcome.frame, config, profile_name)
    if model is not None:
        outcome.model_calls = model.calls
    return outcome


def commit(
    outcome: RunResult, config: Config, root: Path | str | None = None, dry_run: bool = False
) -> dict:
    """Write a run's frames into the warehouse, appending rather than rebuilding."""
    warehouse = Warehouse(config, root, outcome.profile_name or None)
    frames = [r.frame for r in outcome.ok_results]
    sources = [Path(r.meta.path) for r in outcome.ok_results]
    return warehouse.append(frames, sources, dry_run=dry_run)


def profile_headers(root: Path | str, config: Config) -> pd.DataFrame:
    """Inventory every distinct header across every file, with a proposed mapping.

    This is the phase-one deliverable: it is what the target schema gets chosen from,
    and it shows how much of the long tail the model would actually have to handle.
    """
    from openpyxl import load_workbook

    from .detect import extract_table

    mapper = Mapper(config)
    vocabulary = set(mapper.seed) | set(config.aliases.global_)

    seen: dict[str, dict] = {}
    for meta in scan(Path(root), config.banks):
        try:
            workbook = load_workbook(meta.path, data_only=True)
        except Exception:
            continue
        try:
            for sheet_name in workbook.sheetnames:
                table = extract_table(
                    workbook[sheet_name], vocabulary, config.settings.detection
                )
                if not table.rows:
                    continue
                results = mapper.resolve_table(
                    table.headers, table.rows, meta.bank_code, table.header_parts
                )
                for mapping in results:
                    if not mapping.normalized:
                        continue
                    entry = seen.setdefault(mapping.normalized, {
                        "en_tete_normalise": mapping.normalized,
                        "variantes": set(),
                        "banques": set(),
                        "fichiers": 0,
                        "champ_propose": mapping.canonical or "(non résolu)",
                        "methode": mapping.method,
                        "exemples": mapping.profile.samples if mapping.profile else [],
                        "profil": mapping.profile.describe() if mapping.profile else "",
                    })
                    entry["variantes"].add(mapping.header)
                    if meta.bank_code:
                        entry["banques"].add(meta.bank_code)
                    entry["fichiers"] += 1
                break   # the first sheet with data is the sales table
        finally:
            workbook.close()

    rows = []
    for entry in seen.values():
        rows.append({
            "en_tete_normalise": entry["en_tete_normalise"],
            "variantes": " | ".join(sorted(entry["variantes"])),
            "banques": ", ".join(sorted(entry["banques"])),
            "fichiers": entry["fichiers"],
            "champ_propose": entry["champ_propose"],
            "methode": entry["methode"],
            "profil": entry["profil"],
            "exemples": ", ".join(entry["exemples"][:3]),
        })
    frame = pd.DataFrame(rows)
    if frame.empty:
        return frame
    return frame.sort_values(["champ_propose", "fichiers"], ascending=[True, False]).reset_index(drop=True)
